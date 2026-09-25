import os
import re
import time
import queue
import sqlite3
import threading
from flask import Flask, request, jsonify
from flask_cors import CORS
import openpyxl

app = Flask(__name__)
CORS(app)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILENAME = '22.08.2026 genel liste.xlsx'
EXCEL_PATH = os.path.join(BASE_DIR, EXCEL_FILENAME)
DB_PATH = os.path.join(BASE_DIR, 'cache.db')

update_queue = queue.Queue()

def normalize_tr(s):
    if not s:
        return ''
    s = str(s).upper()
    s = s.replace('İ', 'I').replace('ı', 'i').replace('Ş', 'S').replace('ş', 's')
    s = s.replace('Ğ', 'G').replace('ğ', 'g').replace('Ü', 'U').replace('ü', 'u')
    s = s.replace('Ö', 'O').replace('ö', 'o').replace('Ç', 'C').replace('ç', 'c')
    s = s.replace('', '')
    s = s.lower().strip()
    return s

def clean_district(d):
    if not d:
        return ''
    d = str(d).strip().upper()
    if 'GEBZE' in d:
        return 'GEBZE'
    if 'DARICA' in d or 'DAR' in d:
        return 'DARICA'
    if 'AYIROVA' in d or 'CAYIROVA' in d:
        return 'ÇAYIROVA'
    if 'LOVASI' in d or 'DILOVASI' in d:
        return 'DİLOVASI'
    return d

def get_db_connection():
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    return conn

def init_db(force_rebuild=False):
    db_exists = os.path.exists(DB_PATH)
    if not db_exists or force_rebuild:
        print(f"[*] Building database cache from {EXCEL_FILENAME}...")
        t0 = time.time()
        
        if os.path.exists(DB_PATH):
            os.remove(DB_PATH)
            
        conn = get_db_connection()
        c = conn.cursor()
        c.execute('''CREATE TABLE IF NOT EXISTS members (
            row_index INTEGER PRIMARY KEY,
            uye_sicil_no TEXT,
            ticaret_sicil_no TEXT,
            unvan TEXT,
            adres TEXT,
            ilce TEXT,
            clean_ilce TEXT,
            telefon TEXT,
            yetkili TEXT,
            meslek_grubu TEXT,
            yetki_belgesi_durum TEXT,
            durum TEXT,
            belgeyi_alan TEXT,
            referans TEXT,
            notlar TEXT,
            search_text TEXT
        )''')
        
        wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True)
        ws = wb.active
        
        rows_to_insert = []
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 1):
            uye_sicil = str(row[0] or '').strip()
            tic_sicil = str(row[1] or '').strip()
            unvan = str(row[2] or '').strip()
            adres = str(row[3] or '').strip()
            ilce = str(row[4] or '').strip()
            c_ilce = clean_district(ilce)
            tel = str(row[5] or '').strip()
            yetkili = str(row[6] or '').strip()
            meslek = str(row[7] or '').strip()
            yetki_durum = str(row[8] or '').strip() if len(row) > 8 else ''
            durum = str(row[10] or '').strip() if len(row) > 10 else ''
            belgeyi_alan = str(row[11] or '').strip() if len(row) > 11 else ''
            referans = str(row[12] or '').strip() if len(row) > 12 else ''
            notlar = str(row[13] or '').strip() if len(row) > 13 else ''
            
            search_parts = [
                normalize_tr(uye_sicil),
                normalize_tr(tic_sicil),
                normalize_tr(unvan),
                normalize_tr(yetkili),
                normalize_tr(ilce),
                normalize_tr(c_ilce),
                normalize_tr(adres),
                normalize_tr(referans)
            ]
            search_text = ' '.join(search_parts)
            
            rows_to_insert.append((
                idx, uye_sicil, tic_sicil, unvan, adres, ilce, c_ilce, tel, yetkili, meslek, yetki_durum, 
                durum, belgeyi_alan, referans, notlar, search_text
            ))
            
        c.executemany('INSERT INTO members VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)', rows_to_insert)
        c.execute('CREATE INDEX IF NOT EXISTS idx_search ON members(search_text)')
        c.execute('CREATE INDEX IF NOT EXISTS idx_clean_ilce ON members(clean_ilce)')
        c.execute('CREATE INDEX IF NOT EXISTS idx_durum ON members(durum)')
        conn.commit()
        wb.close()
        conn.close()
        print(f"[+] Cache initialized with {len(rows_to_insert)} records in {time.time() - t0:.2f}s")
    else:
        print("[+] Using existing database cache.")

def excel_worker():
    """Background worker that commits status updates to Excel file"""
    pending_updates = {}
    while True:
        try:
            item = update_queue.get(timeout=2)
            row_idx, payload = item
            if row_idx in pending_updates:
                pending_updates[row_idx].update(payload)
            else:
                pending_updates[row_idx] = payload
            update_queue.task_done()
        except queue.Empty:
            pass

        if pending_updates and update_queue.empty():
            items_to_save = dict(pending_updates)
            pending_updates.clear()
            try:
                print(f"[*] Persisting {len(items_to_save)} updates to {EXCEL_FILENAME}...")
                t0 = time.time()
                wb = openpyxl.load_workbook(EXCEL_PATH)
                ws = wb.active
                
                # Column 11: DURUM, Column 12: BELGEYİ ALAN, Column 13: REFERANS, Column 14: NOTLAR
                headers = {
                    11: 'DURUM',
                    12: 'BELGEYİ ALAN',
                    13: 'REFERANS',
                    14: 'NOTLAR'
                }
                for col_idx, h_text in headers.items():
                    if ws.cell(row=1, column=col_idx).value != h_text:
                        ws.cell(row=1, column=col_idx, value=h_text)

                for r_idx, data in items_to_save.items():
                    target_row = r_idx + 1
                    if 'durum' in data:
                        ws.cell(row=target_row, column=11, value=data['durum'] or None)
                    if 'belgeyi_alan' in data:
                        ws.cell(row=target_row, column=12, value=data['belgeyi_alan'] or None)
                    if 'referans' in data:
                        ws.cell(row=target_row, column=13, value=data['referans'] or None)
                    if 'notlar' in data:
                        ws.cell(row=target_row, column=14, value=data['notlar'] or None)

                wb.save(EXCEL_PATH)
                wb.close()
                print(f"[+] Excel saved successfully in {time.time() - t0:.2f}s")
            except Exception as e:
                print(f"[!] Error saving Excel: {e}")

writer_thread = threading.Thread(target=excel_worker, daemon=True)
writer_thread.start()

@app.route('/api/search', methods=['GET'])
def search_members():
    q = request.args.get('q', '').strip()
    page = int(request.args.get('page', 1))
    limit = int(request.args.get('limit', 20))
    ilce = request.args.get('ilce', '').strip()
    durum_filter = request.args.get('durum', '').strip()

    offset = (page - 1) * limit
    normalized_q = normalize_tr(q)
    tokens = [t for t in normalized_q.split() if t]

    conn = get_db_connection()
    c = conn.cursor()

    conditions = []
    params = []

    for token in tokens:
        conditions.append("search_text LIKE ?")
        params.append(f"%{token}%")

    if ilce:
        c_ilce = clean_district(ilce)
        conditions.append("(clean_ilce = ? OR UPPER(ilce) LIKE ?)")
        params.extend([c_ilce, f"%{c_ilce}%"])

    if durum_filter:
        conditions.append("durum = ?")
        params.append(durum_filter)

    where_clause = " AND ".join(conditions) if conditions else "1=1"

    count_sql = f"SELECT COUNT(*) FROM members WHERE {where_clause}"
    c.execute(count_sql, params)
    total_count = c.fetchone()[0]

    select_sql = f"""
        SELECT row_index, uye_sicil_no, ticaret_sicil_no, unvan, adres, ilce, clean_ilce, telefon, yetkili, 
               meslek_grubu, yetki_belgesi_durum, durum, belgeyi_alan, referans, notlar
        FROM members
        WHERE {where_clause}
        ORDER BY row_index ASC
        LIMIT ? OFFSET ?
    """
    c.execute(select_sql, params + [limit, offset])
    rows = c.fetchall()
    conn.close()

    results = []
    for r in rows:
        results.append({
            '__originalRowIndex': r['row_index'],
            'UYE SICIL NO': r['uye_sicil_no'],
            'TİCARET SİCİL NO': r['ticaret_sicil_no'],
            'UNVAN': r['unvan'],
            'ADRES': r['adres'],
            'İLÇE': r['clean_ilce'] or r['ilce'],
            'CEP TELEFONU (GSM) ': r['telefon'],
            'YETKİLİ ADI SOYADI/UNVAN': r['yetkili'],
            'MESLEK GRUBU': r['meslek_grubu'],
            'YETKİ BELGESİ DURUM': r['yetki_belgesi_durum'],
            'DURUM': r['durum'] or '',
            'BELGEYİ ALAN': r['belgeyi_alan'] or '',
            'REFERANS': r['referans'] or '',
            'NOTLAR': r['notlar'] or ''
        })

    return jsonify({
        'results': results,
        'total': total_count,
        'page': page,
        'limit': limit,
        'hasMore': (offset + limit) < total_count
    })

@app.route('/api/update-status', methods=['POST'])
def update_status():
    data = request.get_json() or {}
    row_index = data.get('rowIndex')
    
    if row_index is None:
        return jsonify({'error': 'rowIndex is required'}), 400

    conn = get_db_connection()
    c = conn.cursor()

    # Get current values
    c.execute("SELECT durum, belgeyi_alan, referans, notlar FROM members WHERE row_index = ?", (row_index,))
    curr = c.fetchone()
    if not curr:
        conn.close()
        return jsonify({'error': 'Member not found'}), 404

    durum_val = data.get('value', data.get('durum', curr['durum']))
    belgeyi_alan_val = data.get('belgeyiAlan', curr['belgeyi_alan'])
    referans_val = data.get('referans', curr['referans'])
    notlar_val = data.get('notlar', curr['notlar'])

    c.execute("""
        UPDATE members 
        SET durum = ?, belgeyi_alan = ?, referans = ?, notlar = ? 
        WHERE row_index = ?
    """, (durum_val, belgeyi_alan_val, referans_val, notlar_val, row_index))
    conn.commit()
    conn.close()

    payload = {
        'durum': durum_val,
        'belgeyi_alan': belgeyi_alan_val,
        'referans': referans_val,
        'notlar': notlar_val
    }
    update_queue.put((int(row_index), payload))

    return jsonify({
        'success': True,
        'rowIndex': row_index,
        'durum': durum_val,
        'belgeyiAlan': belgeyi_alan_val,
        'referans': referans_val,
        'notlar': notlar_val
    })

@app.route('/api/stats', methods=['GET'])
def get_stats():
    conn = get_db_connection()
    c = conn.cursor()

    c.execute("SELECT COUNT(*) FROM members")
    total = c.fetchone()[0]

    c.execute("SELECT COUNT(*) FROM members WHERE durum = 'Ziyarete Gidildi'")
    visited = c.fetchone()[0]

    c.execute("SELECT COUNT(*) FROM members WHERE durum = 'Belge Alındı'")
    document_received = c.fetchone()[0]

    c.execute("SELECT COUNT(*) FROM members WHERE durum = 'Karşı Tarafta'")
    opposite_party = c.fetchone()[0]

    c.execute("""
        SELECT clean_ilce, COUNT(*) as count 
        FROM members 
        WHERE clean_ilce IN ('GEBZE', 'DARICA', 'ÇAYIROVA', 'DİLOVASI')
        GROUP BY clean_ilce 
        ORDER BY count DESC
    """)
    districts = [{'ilce': r['clean_ilce'], 'count': r['count']} for r in c.fetchall()]

    conn.close()

    return jsonify({
        'total': total,
        'visited': visited,
        'document_received': document_received,
        'opposite_party': opposite_party,
        'districts': districts
    })

# Serve static frontend if built
DIST_DIR = os.path.join(BASE_DIR, 'dist')

@app.route('/', defaults={'path': ''})
@app.route('/<path:path>')
def serve_frontend(path):
    if path.startswith('api/'):
        return jsonify({'error': 'Not found'}), 404
        
    full_path = os.path.join(DIST_DIR, path)
    if path and os.path.exists(full_path):
        from flask import send_from_directory
        return send_from_directory(DIST_DIR, path)
        
    index_file = os.path.join(DIST_DIR, 'index.html')
    if os.path.exists(index_file):
        from flask import send_file
        return send_file(index_file)
        
    return "Lütfen 'npm run build' çalıştırın veya Vite dev sunucusunu kullanın.", 200

if __name__ == '__main__':
    init_db(force_rebuild=False)
    print("[*] Starting backend server on http://localhost:3001 ...")
    app.run(host='0.0.0.0', port=3001, debug=False)
